#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""把 docs/玩家手册.md 转为 docs/玩家手册.docx（中文友好的格式）。"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from pathlib import Path
import re


SRC = Path("docs/玩家手册.md")
DST = Path("docs/玩家手册.docx")


def add_para(doc, text, *, bold=False, italic=False, size=11, align=None, color=None, space_after=4):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph()
    if level == 1:
        size = 22
    elif level == 2:
        size = 16
    else:
        size = 13
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = True
    p.paragraph_format.space_before = Pt(10 if level <= 2 else 6)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    # 段落里可能含 inline bold（**xxx**）
    for chunk, is_bold in _split_bold(text):
        r = p.add_run(chunk)
        r.bold = is_bold
        r.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(2)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    for chunk, is_bold in _split_bold(text):
        r = p.add_run(chunk)
        r.bold = is_bold
        r.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(2)
    return p


def add_quote(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.right_indent = Cm(0.4)
    run = p.add_run("「" + text + "」")
    run.font.size = Pt(11)
    run.italic = True
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_tip(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.right_indent = Cm(0.3)
    r1 = p.add_run("💡 ")
    r1.font.size = Pt(11)
    r1.bold = True
    r2 = p.add_run(text)
    r2.font.size = Pt(11)
    r2.italic = True
    r2.font.color.rgb = RGBColor(0x33, 0x66, 0x33)
    p.paragraph_format.space_after = Pt(4)
    return p


def _split_bold(text):
    """把 **xxx** 切出来，返回 [(text, is_bold), ...]"""
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    out = []
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            out.append((part[2:-2], True))
        else:
            out.append((part, False))
    return out


def add_runs(p, text, *, size=11):
    for chunk, is_bold in _split_bold(text):
        r = p.add_run(chunk)
        r.bold = is_bold
        r.font.size = Pt(size)


def parse_and_build():
    doc = Document()
    # 全局字体 / 边距
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(11)
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    text = SRC.read_text(encoding="utf-8")
    lines = text.split("\n")

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        # 标题
        if line.startswith("# "):
            add_heading(doc, line[2:].strip(), 1)
        elif line.startswith("## "):
            add_heading(doc, line[3:].strip(), 2)
        elif line.startswith("### "):
            add_heading(doc, line[4:].strip(), 3)
        # 引用块（> ...）
        elif line.startswith("> "):
            add_quote(doc, line[2:].strip())
        # 无序列表
        elif line.startswith("- "):
            add_bullet(doc, line[2:].strip())
        # 数字列表
        elif re.match(r"^\d+\.\s", line):
            add_numbered(doc, re.sub(r"^\d+\.\s", "", line))
        # 普通段落
        else:
            p = doc.add_paragraph()
            add_runs(p, line)
            p.paragraph_format.space_after = Pt(4)
        i += 1

    doc.save(DST)
    print(f"OK: {DST} ({DST.stat().st_size} bytes)")


if __name__ == "__main__":
    parse_and_build()