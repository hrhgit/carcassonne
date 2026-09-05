"""把 PROGRAM_DOC.md 转成 DOCX,保留标题/列表/表格/代码块结构"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import re
import sys

SRC = r"C:\Users\Luren_jybd\WorkBuddy\2026-09-05-15-10-40\carcassonne\docs\PROGRAM_DOC.md"
DST = r"C:\Users\Luren_jybd\WorkBuddy\2026-09-05-15-10-40\carcassonne\docs\PROGRAM_DOC.docx"

# 样式颜色（碧水沃野主题色）
THEME_GREEN = RGBColor(0x2E, 0x7D, 0x32)
THEME_DARK = RGBColor(0x1A, 0x2E, 0x1A)
TEXT_DARK = RGBColor(0x22, 0x33, 0x22)


def set_cell_shading(cell, hex_color: str):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def add_inline(paragraph, text):
    """处理加粗、行内代码、斜体"""
    # 简单处理：**bold** 和 `code` 和 *italic*
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
            run.font.color.rgb = THEME_GREEN
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
        run.font.size = Pt(11)


def main():
    with open(SRC, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()

    # 全局样式
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(11)
    style.font.color.rgb = TEXT_DARK

    # 页面边距
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    # 标题样式
    for level, size in [(1, 24), (2, 18), (3, 14), (4, 12)]:
        s = doc.styles[f"Heading {level}"]
        s.font.name = "Microsoft YaHei"
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = THEME_GREEN if level == 1 else THEME_DARK

    i = 0
    while i < len(lines):
        line = lines[i].rstrip("\n")

        # 空行
        if not line.strip():
            i += 1
            continue

        # 水平线
        if line.strip() == "---":
            p = doc.add_paragraph()
            p.add_run("─" * 60).font.color.rgb = RGBColor(0xAA, 0xBB, 0xAA)
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            title = m.group(2).strip()
            p = doc.add_heading(title, level=level)
            if level == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # 代码块
        if line.strip().startswith("```"):
            lang = line.strip()[3:].strip()
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i].rstrip("\n"))
                i += 1
            i += 1  # 跳过闭合 ```
            p = doc.add_paragraph()
            for j, cl in enumerate(code_lines):
                if j > 0:
                    p.add_run("\n").font.size = Pt(10)
                run = p.add_run(cl)
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x33, 0x55, 0x33)
                # 灰色背景用 shading on paragraph
            p.paragraph_format.left_indent = Cm(0.6)
            from docx.oxml.ns import qn
            from docx.oxml import OxmlElement
            pPr = p._p.get_or_add_pPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'), 'F1F5EE')
            pPr.append(shd)
            continue

        # 表格
        if line.lstrip().startswith("|") and i + 1 < len(lines) \
                and re.match(r"^\s*\|[\s\-:|]+\|\s*$", lines[i + 1]):
            header_cells = [c.strip() for c in line.strip().strip("|").split("|")]
            i += 2  # 跳过分隔行
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            table = doc.add_table(rows=1 + len(rows), cols=len(header_cells))
            table.style = "Light Grid Accent 1"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, h in enumerate(header_cells):
                cell = table.rows[0].cells[j]
                cell.text = ""
                p = cell.paragraphs[0]
                run = p.add_run(h)
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.size = Pt(10)
                set_cell_shading(cell, "2E7D32")
            for r_idx, row in enumerate(rows):
                for c_idx, val in enumerate(row):
                    if c_idx >= len(header_cells):
                        break
                    cell = table.rows[r_idx + 1].cells[c_idx]
                    cell.text = ""
                    p = cell.paragraphs[0]
                    add_inline(p, val)
                    for run in p.runs:
                        if not run.font.size:
                            run.font.size = Pt(10)
            doc.add_paragraph()  # 表格后空一行
            continue

        # 引用块
        if line.startswith(">"):
            text = line.lstrip(">").strip()
            p = doc.add_paragraph()
            run = p.add_run("▌ " + text)
            run.font.color.rgb = THEME_GREEN
            run.italic = True
            run.font.size = Pt(11)
            p.paragraph_format.left_indent = Cm(0.5)
            i += 1
            continue

        # 列表项
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if m:
            indent = len(m.group(1)) // 2
            marker = m.group(2)
            text = m.group(3)
            style = "List Bullet" if marker in ("-", "*") else "List Number"
            p = doc.add_paragraph(style=style)
            add_inline(p, text)
            p.paragraph_format.left_indent = Cm(0.6 + indent * 0.6)
            for run in p.runs:
                if not run.font.size:
                    run.font.size = Pt(11)
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_inline(p, line)
        for run in p.runs:
            if not run.font.size:
                run.font.size = Pt(11)

        i += 1

    doc.save(DST)
    print(f"DOCX 已写入: {DST}")


if __name__ == "__main__":
    main()